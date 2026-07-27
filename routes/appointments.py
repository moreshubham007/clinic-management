from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from flask_login import login_required, current_user
from extensions import db
from models import User, Doctor, Appointment, Case, CaseHistory, WaitingArea, AppointmentRequest
from datetime import datetime, timedelta
from functools import wraps
import io

appointments_bp = Blueprint('appointments', __name__)

def doctor_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'doctor':
            flash('You need to be a doctor to access this page.', 'error')
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated_function

def receptionist_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if not current_user.is_authenticated or current_user.role != 'receptionist':
            flash('You need to be a receptionist to access this page.', 'error')
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated_function

@appointments_bp.route('/')
@login_required
def list_appointments():
    # Base query
    query = Appointment.query

    # Apply role-based filters
    if current_user.role == 'doctor':
        query = query.filter_by(doctor_id=current_user.doctor.id)
    elif current_user.role == 'patient':
        query = query.filter_by(patient_id=current_user.id)
    elif current_user.role not in ['admin', 'receptionist']:
        flash('You do not have permission to view appointments', 'danger')
        return redirect(url_for('index'))
    
    # Apply status filter
    status = request.args.get('status')
    if status:
        query = query.filter_by(status=status)
    
    # Apply date filter
    date = request.args.get('date')
    if date:
        try:
            filter_date = datetime.strptime(date, '%Y-%m-%d').date()
            query = query.filter(db.func.date(Appointment.datetime) == filter_date)
        except ValueError:
            flash('Invalid date format', 'warning')
    
    # Apply doctor filter (for admin/receptionist)
    doctor_id = request.args.get('doctor')
    if doctor_id and current_user.role in ['admin', 'receptionist']:
        try:
            query = query.filter_by(doctor_id=int(doctor_id))
        except ValueError:
            flash('Invalid doctor selection', 'warning')
    
    # Add patient_type filter
    if request.args.get('patient_type'):
        query = query.filter(Appointment.patient_type == request.args.get('patient_type'))

    # Add priority filter
    if request.args.get('priority'):
        query = query.filter(Appointment.priority == request.args.get('priority'))

    # Add payment_status filter
    if request.args.get('payment_status'):
        query = query.filter(Appointment.payment_status == request.args.get('payment_status'))

    # Add has_bill filter
    has_bill = request.args.get('has_bill')
    if has_bill == 'yes':
        query = query.filter(
            db.or_(Appointment.consultation_fee > 0, Appointment.medicine_charges > 0)
        )
    elif has_bill == 'no':
        query = query.filter(
            db.and_(
                db.or_(Appointment.consultation_fee == None, Appointment.consultation_fee == 0),
                db.or_(Appointment.medicine_charges == None, Appointment.medicine_charges == 0)
            )
        )

    # Get all active doctors for the filter dropdown
    doctors = Doctor.query.join(User).filter(User.is_active == True).all()

    page = request.args.get('page', 1, type=int)
    per_page = 20

    # Order by datetime in descending order (newest first)
    pagination = query.order_by(Appointment.datetime.desc()).paginate(
        page=page, per_page=per_page, error_out=False
    )

    return render_template('appointments/list.html',
                         appointments=pagination.items,
                         pagination=pagination,
                         doctors=doctors)

@appointments_bp.route('/create', methods=['GET', 'POST'])
@login_required
def create_appointment():
    if current_user.role not in ['admin', 'receptionist', 'patient', 'doctor']:
        flash('You do not have permission to create appointments', 'danger')
        return redirect(url_for('appointments.list_appointments'))
    
    # Preserve req_id and patient_id across validation failures
    req_id_param = request.args.get('req_id', type=int)
    patient_id_param = request.args.get('patient_id', type=int)

    if request.method == 'POST':
        # For doctors, automatically set themselves as the doctor
        if current_user.role == 'doctor':
            doctor_id = current_user.doctor.id
        else:
            doctor_id = request.form.get('doctor_id')

        # Accept patient_id from form body OR query string
        if current_user.role in ['admin', 'receptionist', 'doctor']:
            patient_id = request.form.get('patient_id') or request.args.get('patient_id')
        else:
            patient_id = current_user.id

        date_str = request.form.get('date')
        time_str = request.form.get('time')

        if not all([doctor_id, patient_id, date_str, time_str]):
            flash('Please fill in all required fields', 'danger')
            return redirect(url_for('appointments.create_appointment',
                                    req_id=req_id_param,
                                    patient_id=patient_id or patient_id_param))
        
        try:
            appointment_datetime = datetime.strptime(f"{date_str} {time_str}", "%Y-%m-%d %H:%M")
            
            # Check if the appointment time is in the future
            if appointment_datetime < datetime.now():
                flash('Appointment time must be in the future', 'danger')
                return redirect(url_for('appointments.create_appointment',
                                        req_id=req_id_param, patient_id=patient_id))

            # Check if the doctor is available at this time
            existing_appointment = Appointment.query.filter_by(
                doctor_id=doctor_id,
                datetime=appointment_datetime
            ).first()

            if existing_appointment:
                flash('This time slot is already booked', 'danger')
                return redirect(url_for('appointments.create_appointment',
                                        req_id=req_id_param, patient_id=patient_id))
            
            appointment = Appointment(
                doctor_id=doctor_id,
                patient_id=patient_id,
                datetime=appointment_datetime,
                patient_type=request.form.get('patient_type', 'existing'),
                priority=request.form.get('priority', 'medium'),
                notes=request.form.get('notes'),
                remarks=request.form.get('remarks') if current_user.role == 'doctor' else None
            )
            
            db.session.add(appointment)
            db.session.commit()
            
            flash('Appointment created successfully', 'success')
            return redirect(url_for('appointments.list_appointments'))
            
        except ValueError:
            flash('Invalid date or time format', 'danger')
            return redirect(url_for('appointments.create_appointment',
                                    req_id=req_id_param, patient_id=patient_id_param))

    # For GET request - prepare data for the form
    if current_user.role == 'doctor':
        doctors = [current_user.doctor]
    else:
        doctors = Doctor.query.join(User).filter(User.is_active == True).all()

    patients = User.query.filter_by(role='patient').all() if current_user.role in ['admin', 'receptionist', 'doctor'] else None

    # Prefill patient from req_id or direct patient_id query param
    prefill_patient = None
    prefill_req = None

    if req_id_param:
        prefill_req = AppointmentRequest.query.get(req_id_param)
        if prefill_req and prefill_req.patient_id:
            prefill_patient = User.query.get(prefill_req.patient_id)

    # Fallback: patient_id passed directly in query string
    if not prefill_patient and patient_id_param:
        prefill_patient = User.query.get(patient_id_param)

    return render_template('appointments/create.html',
                           doctors=doctors,
                           patients=patients,
                           prefill_patient=prefill_patient,
                           prefill_req=prefill_req,
                           req_id_param=req_id_param)

@appointments_bp.route('/<int:appointment_id>/cancel', methods=['POST'])
@login_required
def cancel_appointment(appointment_id):
    appointment = Appointment.query.get_or_404(appointment_id)
    
    # Check permissions
    if current_user.role not in ['admin', 'receptionist', 'doctor'] and current_user.id != appointment.patient_id:
        return jsonify({'error': 'Unauthorized'}), 403
    
    appointment.status = 'cancelled'
    db.session.commit()
    
    return jsonify({'success': True})

@appointments_bp.route('/<int:appointment_id>/complete', methods=['POST'])
@login_required
def complete_appointment(appointment_id):
    if current_user.role not in ['admin', 'receptionist', 'doctor']:
        return jsonify({'error': 'Unauthorized'}), 403
    
    appointment = Appointment.query.get_or_404(appointment_id)
    appointment.status = 'completed'
    
    # Update any waiting area entry for this appointment
    waiting_entry = WaitingArea.query.filter_by(appointment_id=appointment_id).first()
    if waiting_entry and waiting_entry.status in ['waiting', 'in_progress']:
        waiting_entry.status = 'completed'
        waiting_entry.completion_time = datetime.now()
        waiting_entry.updated_at = datetime.now()
    
    db.session.commit()
    
    return jsonify({'success': True})

@appointments_bp.route('/doctor/<int:doctor_id>/availability')
def get_doctor_availability(doctor_id):
    doctor = Doctor.query.get_or_404(doctor_id)
    date_str = request.args.get('date')
    
    try:
        date = datetime.strptime(date_str, "%Y-%m-%d").date()
        
        # Get all appointments for the doctor on this date
        appointments = Appointment.query.filter(
            Appointment.doctor_id == doctor_id,
            db.func.date(Appointment.datetime) == date,
            Appointment.status != 'cancelled'
        ).all()
        
        # Convert appointments to time slots
        booked_slots = [appt.datetime.strftime("%H:%M") for appt in appointments]
        
        # Get doctor's availability for this day of week
        day_of_week = date.strftime("%A").lower()
        availability = doctor.availability.get(day_of_week, [])
        
        return jsonify({
            'available_slots': availability,
            'booked_slots': booked_slots
        })
        
    except ValueError:
        return jsonify({'error': 'Invalid date format'}), 400

@appointments_bp.route('/cases/<int:patient_id>')
@doctor_required
def view_patient_cases(patient_id):
    cases = Case.query.filter_by(patient_id=patient_id).all()
    return render_template('cases/list.html', cases=cases)

@appointments_bp.route('/cases/create/<int:appointment_id>', methods=['GET', 'POST'])
@doctor_required
def create_case(appointment_id):
    appointment = Appointment.query.get_or_404(appointment_id)
    doctor = Doctor.query.filter_by(user_id=current_user.id).first()
    
    if request.method == 'POST':
        diagnosis = request.form.get('diagnosis')
        treatment = request.form.get('treatment')
        show_to_patient = request.form.get('show_to_patient') == 'on'
        
        case = Case(
            patient_id=appointment.patient_id,
            doctor_id=doctor.id,
            diagnosis=diagnosis,
            treatment=treatment,
            show_to_patient=show_to_patient
        )
        db.session.add(case)
        
        case_history = CaseHistory(
            case_id=case.id,
            doctor_id=doctor.id,
            notes=f"Initial case created by Dr. {current_user.name}"
        )
        db.session.add(case_history)
        db.session.commit()
        
        flash('Case created successfully', 'success')
        return redirect(url_for('appointments.view_patient_cases', patient_id=appointment.patient_id))
    
    return render_template('cases/create.html', appointment=appointment)

@appointments_bp.route('/cases/<int:case_id>/transfer', methods=['POST'])
@doctor_required
def transfer_case(case_id):
    case = Case.query.get_or_404(case_id)
    new_doctor_id = request.form.get('new_doctor_id')
    transfer_notes = request.form.get('transfer_notes')
    
    case.doctor_id = new_doctor_id
    
    case_history = CaseHistory(
        case_id=case.id,
        doctor_id=current_user.id,
        notes=f"Case transferred to Dr. {User.query.get(new_doctor_id).name}. Notes: {transfer_notes}"
    )
    db.session.add(case_history)
    db.session.commit()
    
    flash('Case transferred successfully', 'success')
    return redirect(url_for('appointments.view_patient_cases', patient_id=case.patient_id))

@appointments_bp.route('/<int:appointment_id>/edit', methods=['GET', 'POST'])
@login_required
def update_appointment(appointment_id):
    appointment = Appointment.query.get_or_404(appointment_id)
    
    # Enhanced permission checking based on appointment status
    can_edit = False
    
    if current_user.role == 'admin':
        can_edit = True
    elif appointment.status == 'completed':
        # Only doctors can edit completed appointments (besides admin)
        if current_user.role == 'doctor' and current_user.doctor.id == appointment.doctor_id:
            can_edit = True
    else:
        # Non-completed appointments: admin, receptionist, or assigned doctor
        if current_user.role in ['admin', 'receptionist'] or \
           (current_user.role == 'doctor' and current_user.doctor.id == appointment.doctor_id):
            can_edit = True
    
    if not can_edit:
        if appointment.status == 'completed' and current_user.role == 'receptionist':
            flash('Completed appointments can only be edited by the assigned doctor or admin', 'warning')
        else:
            flash('You do not have permission to edit this appointment', 'danger')
        return redirect(url_for('appointments.list_appointments'))
    
    if request.method == 'POST':
        try:
            # Update appointment details
            appointment.datetime = datetime.combine(
                datetime.strptime(request.form['date'], '%Y-%m-%d').date(),
                datetime.strptime(request.form['time'], '%H:%M').time()
            )
            appointment.doctor_id = request.form['doctor_id']
            appointment.patient_type = request.form.get('patient_type', 'existing')
            appointment.priority = request.form.get('priority', 'medium')
            appointment.notes = request.form.get('notes', '')
            
            # Allow admin/receptionist to change patient and status (if not completed)
            if current_user.role in ['admin', 'receptionist']:
                # Receptionist can only modify non-completed appointments
                if appointment.status != 'completed' or current_user.role == 'admin':
                    appointment.patient_id = request.form.get('patient_id', appointment.patient_id)
                    appointment.status = request.form.get('status', appointment.status)
            
            # Allow doctor to add remarks
            if current_user.role == 'doctor':
                appointment.remarks = request.form.get('remarks', '')
            
            db.session.commit()
            flash('Appointment updated successfully', 'success')
            return redirect(url_for('appointments.list_appointments'))
            
        except ValueError:
            flash('Invalid date or time format', 'danger')
            return redirect(url_for('appointments.update_appointment', appointment_id=appointment_id))
    
    # Get all active doctors and patients for dropdowns
    doctors = Doctor.query.join(User).filter(User.is_active == True).all()
    patients = User.query.filter(User.role == 'patient', User.is_active == True).all()
    
    return render_template('appointments/edit.html', 
                         appointment=appointment,
                         doctors=doctors,
                         patients=patients,
                         now=datetime.now())

@appointments_bp.route('/<int:appointment_id>/delete', methods=['POST'])
@login_required
def delete_appointment(appointment_id):
    # Check if user is a receptionist
    if current_user.role != 'receptionist':
        flash('You do not have permission to delete appointments', 'danger')
        return redirect(url_for('appointments.list_appointments'))
    
    appointment = Appointment.query.get_or_404(appointment_id)
    
    # Only allow deletion of scheduled appointments
    if appointment.status != 'scheduled':
        flash('Only scheduled appointments can be deleted', 'danger')
        return redirect(url_for('appointments.list_appointments'))
    
    try:
        db.session.delete(appointment)
        db.session.commit()
        flash('Appointment deleted successfully', 'success')
    except Exception as e:
        db.session.rollback()
        flash('Error deleting appointment', 'danger')
    
    return redirect(url_for('appointments.list_appointments'))

@appointments_bp.route('/search-patient')
@login_required
def search_patient():
    search = request.args.get('q', '')
    if not search:
        return jsonify([])
    
    search_term = f"%{search}%"
    patients = User.query.filter(
        User.role == 'patient',
        (User.name.like(search_term)) |
        (User.mobile_number.like(search_term)) |
        (User.patient_number.like(search_term))
    ).limit(10).all()
    
    return jsonify([{
        'id': patient.id,
        'name': patient.name,
        'email': patient.email,
        'mobile_number': patient.mobile_number,
        'patient_number': patient.patient_number
    } for patient in patients])


def _build_appointment_query():
    """Return a filtered Appointment query based on current request args + user role."""
    query = Appointment.query
    if current_user.role == 'doctor':
        query = query.filter_by(doctor_id=current_user.doctor.id)
    elif current_user.role == 'patient':
        query = query.filter_by(patient_id=current_user.id)

    status = request.args.get('status')
    if status:
        query = query.filter_by(status=status)

    date = request.args.get('date')
    if date:
        try:
            fd = datetime.strptime(date, '%Y-%m-%d').date()
            query = query.filter(db.func.date(Appointment.datetime) == fd)
        except ValueError:
            pass

    doctor_id = request.args.get('doctor')
    if doctor_id and current_user.role in ['admin', 'receptionist']:
        try:
            query = query.filter_by(doctor_id=int(doctor_id))
        except ValueError:
            pass

    patient_type = request.args.get('patient_type')
    if patient_type:
        query = query.filter(Appointment.patient_type == patient_type)

    priority = request.args.get('priority')
    if priority:
        query = query.filter(Appointment.priority == priority)

    payment_status = request.args.get('payment_status')
    if payment_status:
        query = query.filter(Appointment.payment_status == payment_status)

    has_bill = request.args.get('has_bill')
    if has_bill == 'yes':
        query = query.filter(
            db.or_(
                Appointment.consultation_fee > 0,
                Appointment.medicine_charges > 0
            )
        )
    elif has_bill == 'no':
        query = query.filter(
            db.and_(
                db.or_(Appointment.consultation_fee == None, Appointment.consultation_fee == 0),
                db.or_(Appointment.medicine_charges == None, Appointment.medicine_charges == 0)
            )
        )

    return query.order_by(Appointment.datetime.desc())


@appointments_bp.route('/export')
@login_required
def export_appointments():
    if current_user.role not in ['admin', 'doctor', 'receptionist']:
        flash('Permission denied.', 'danger')
        return redirect(url_for('appointments.list_appointments'))

    fmt = request.args.get('format', 'excel')
    appointments = _build_appointment_query().all()

    # ── Print view ────────────────────────────────────────────────────────────
    if fmt == 'print':
        filters = {
            'status':       request.args.get('status', ''),
            'date':         request.args.get('date', ''),
            'doctor':       request.args.get('doctor', ''),
            'patient_type': request.args.get('patient_type', ''),
            'priority':     request.args.get('priority', ''),
        }
        return render_template('appointments/print.html',
                               appointments=appointments,
                               filters=filters,
                               generated_at=datetime.now())

    # ── Excel export ──────────────────────────────────────────────────────────
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash('openpyxl not installed. Please rebuild the Docker image.', 'danger')
        return redirect(url_for('appointments.list_appointments'))

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Appointments'

    # Header style
    hdr_font  = Font(bold=True, color='FFFFFF', size=11)
    hdr_fill  = PatternFill('solid', fgColor='2E7D32')
    thin      = Side(style='thin', color='CCCCCC')
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)
    center    = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = [
        'Date', 'Time', 'Patient Name', 'Patient ID', 'Patient Type',
        'Priority', 'Doctor', 'Status',
        'Consultation (₹)', 'Medicine (₹)', 'Discount (₹)', 'Total Bill (₹)',
        'Payment Status', 'Amount Paid (₹)', 'Payment Mode', 'Notes'
    ]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = center
        cell.border    = border

    # Data rows
    alt_fill = PatternFill('solid', fgColor='F1F8E9')
    for row_num, appt in enumerate(appointments, 2):
        consult = float(appt.consultation_fee or 0)
        meds    = float(appt.medicine_charges  or 0)
        disc    = float(appt.discount          or 0)
        total   = consult + meds - disc

        row = [
            appt.datetime.strftime('%d/%m/%Y'),
            appt.datetime.strftime('%H:%M'),
            appt.patient.name if appt.patient else '',
            appt.patient.patient_number if appt.patient else '',
            (appt.patient_type or '').title(),
            (appt.priority or '').title(),
            f"Dr. {appt.doctor.user.name}" if appt.doctor else '',
            (appt.status or '').title(),
            consult if consult else '',
            meds    if meds    else '',
            disc    if disc    else '',
            total   if total   else '',
            (appt.payment_status or '').title(),
            float(appt.payment_amount) if appt.payment_amount else '',
            (appt.payment_mode or '').title(),
            appt.notes or '',
        ]
        ws.append(row)
        fill = alt_fill if row_num % 2 == 0 else None
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_num, column=col_idx)
            cell.border    = border
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            if fill:
                cell.fill = fill

    # Column widths
    col_widths = [12, 8, 22, 12, 14, 10, 22, 12, 16, 14, 12, 14, 14, 16, 14, 30]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30

    # Freeze header row
    ws.freeze_panes = 'A2'

    # Add summary row at the bottom
    ws.append([])
    summary_row = ws.max_row + 1
    ws.cell(summary_row, 1, f'Total appointments: {len(appointments)}').font = Font(bold=True)
    ws.cell(summary_row, 12, sum(
        float(a.consultation_fee or 0) + float(a.medicine_charges or 0) - float(a.discount or 0)
        for a in appointments
    )).font = Font(bold=True, color='2E7D32')

    # Save to buffer and send
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    date_str = datetime.now().strftime('%Y-%m-%d')
    filename = f'appointments_{date_str}.xlsx'
    return send_file(buf, as_attachment=True,
                     download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@appointments_bp.route('/<int:appointment_id>/billing', methods=['POST'])
@login_required
def update_billing(appointment_id):
    """Doctor or Admin can set consultation fee, medicine charges, discount."""
    if current_user.role not in ['admin', 'doctor']:
        flash('Only doctors and admins can edit billing details.', 'danger')
        return redirect(url_for('appointments.list_appointments'))

    appointment = Appointment.query.get_or_404(appointment_id)

    if current_user.role == 'doctor' and appointment.doctor_id != current_user.doctor.id:
        flash('You can only edit billing for your own appointments.', 'danger')
        return redirect(url_for('appointments.list_appointments'))

    def _parse(val):
        try:
            v = float(val)
            return v if v >= 0 else None
        except (TypeError, ValueError):
            return None

    appointment.consultation_fee  = _parse(request.form.get('consultation_fee'))
    appointment.medicine_charges  = _parse(request.form.get('medicine_charges'))
    appointment.discount          = _parse(request.form.get('discount')) or 0
    appointment.updated_at        = datetime.utcnow()
    db.session.commit()
    flash('Billing details updated.', 'success')
    return redirect(url_for('appointments.list_appointments',
                            status=request.args.get('status', ''),
                            date=request.args.get('date', ''),
                            priority=request.args.get('priority', '')))


@appointments_bp.route('/<int:appointment_id>/payment-status', methods=['POST'])
@login_required
def update_payment_status(appointment_id):
    """Receptionist, Doctor, or Admin can record payment collection."""
    if current_user.role not in ['admin', 'doctor', 'receptionist']:
        flash('Permission denied.', 'danger')
        return redirect(url_for('appointments.list_appointments'))

    appointment = Appointment.query.get_or_404(appointment_id)

    payment_status = request.form.get('payment_status', 'unpaid')
    if payment_status not in ['paid', 'unpaid', 'partial']:
        payment_status = 'unpaid'

    amount_str  = request.form.get('payment_amount', '').strip()
    payment_mode = request.form.get('payment_mode', '').strip() or None

    appointment.payment_status = payment_status
    appointment.payment_mode   = payment_mode

    if amount_str:
        try:
            appointment.payment_amount = float(amount_str)
        except ValueError:
            flash('Invalid payment amount.', 'danger')
            return redirect(url_for('appointments.list_appointments'))
    else:
        appointment.payment_amount = None

    if payment_status in ['paid', 'partial']:
        appointment.payment_received_by = current_user.id
        appointment.payment_date        = datetime.utcnow()
    else:
        appointment.payment_received_by = None
        appointment.payment_date        = None

    appointment.updated_at = datetime.utcnow()
    db.session.commit()
    flash('Payment status updated.', 'success')
    return redirect(url_for('appointments.list_appointments',
                            status=request.args.get('status', ''),
                            date=request.args.get('date', ''),
                            priority=request.args.get('priority', '')))


@appointments_bp.route('/<int:appointment_id>/priority', methods=['POST'])
@login_required
def update_priority(appointment_id):
    if current_user.role not in ['admin', 'doctor']:
        return jsonify({'error': 'Permission denied'}), 403

    appointment = Appointment.query.get_or_404(appointment_id)

    # Doctors can only update priority for their own appointments
    if current_user.role == 'doctor' and appointment.doctor_id != current_user.doctor.id:
        return jsonify({'error': 'Permission denied'}), 403

    data = request.get_json(silent=True) or {}
    new_priority = data.get('priority', '').strip()

    if new_priority not in ['high', 'medium', 'low']:
        return jsonify({'error': 'Invalid priority'}), 400

    appointment.priority = new_priority
    db.session.commit()
    return jsonify({'success': True, 'priority': new_priority})